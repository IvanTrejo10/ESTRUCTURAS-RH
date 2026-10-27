from io import BytesIO
from pathlib import Path
import sys
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from structures import StructureError, bundle, generate


CR_HEADER = ["Ruta", "Territorio", "Subdireccion", "Zona ", "Sucursal", "Marca", "Base", "Unidad de Negocio", "País"]
VALES_HEADER = ["Coordinacion", "Marca", "Dir", "Subdirección", "Zona", "Sucursal", "Red", "Cartera", "Estado"]


def book_bytes(headers, rows, name="Estructuras"):
    book = Workbook()
    sheet = book.active
    sheet.title = name
    for row_number, row in enumerate([headers, *rows], 1):
        for col, value in enumerate(row, 1):
            cell = sheet.cell(row_number, col, value)
            if isinstance(value, str):
                cell.data_type = "s"
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def cr_row(route="RUTA-A", country="México", territory="Territorio A", sub="Subdirección A", zone="Zona A"):
    return [route, territory, sub, zone, "Sucursal A", "Marca A", "Base A", "Negocio A", country]


def vales_row(route="Coordinación A", brand="Vale Amigo"):
    return [route, brand, "Vales", "Subdirección A", "Zona A", "Sucursal A", brand, brand, "Estado A"]


class StructuresTests(unittest.TestCase):
    def test_cr_mexico_preserves_both_fields_and_latam_uses_subdireccion(self):
        data = book_bytes(CR_HEADER, [cr_row(), cr_row("RUTA-B", "Perú", territory="NO APLICA", sub="Sur")])
        result = generate("cr", data, 43, 2026)
        self.assertEqual(result.total, 2)
        self.assertEqual(result.sheets["MEXICO"][0][1:3], ["Territorio A", "Subdirección A"])
        self.assertEqual(result.sheets["PERU"][0][1], "Sur")
        self.assertEqual(len(result.headers["MEXICO"]), 7)
        self.assertEqual(len(result.headers["PERU"]), 6)

    def test_vales_grouping_column_order_and_all_tabs(self):
        data = book_bytes(VALES_HEADER, [vales_row("A", "Vale Amigo"), vales_row("B", "RapiVale")], "Hoja 1")
        result = generate("vales", data, 43, 2026)
        self.assertEqual(result.sheets["Rapi"][0], ["B", "Subdirección A", "Zona A", "Sucursal A", "RapiVale"])
        self.assertEqual(list(result.sheets), ["Rapi", "Peru", "Vale Amigo", "Viva Vale"])
        self.assertEqual(result.sheets["Peru"], [])

    def test_header_accents_and_space_are_normalized(self):
        headers = CR_HEADER.copy()
        headers[2], headers[8] = " SUBDIRECCIÓN ", "PAIS"
        result = generate("cr", book_bytes(headers, [cr_row(country="MEXICO")]), 43, 2026)
        self.assertEqual(result.total, 1)

    def test_duplicate_keys_do_not_silently_drop_data(self):
        with self.assertRaisesRegex(StructureError, "clave repetida"):
            generate("cr", book_bytes(CR_HEADER, [cr_row(), cr_row()]), 43, 2026)

    def test_unknown_country_cannot_disappear(self):
        with self.assertRaisesRegex(StructureError, "no contemplado"):
            generate("cr", book_bytes(CR_HEADER, [cr_row(country="Chile")]), 43, 2026)

    def test_unknown_brand_cannot_disappear(self):
        with self.assertRaisesRegex(StructureError, "no contemplado"):
            generate("vales", book_bytes(VALES_HEADER, [vales_row(brand="Nueva Marca")]), 43, 2026)

    def test_incomplete_record_is_rejected(self):
        row = cr_row()
        row[4] = None
        with self.assertRaisesRegex(StructureError, "falta Sucursal"):
            generate("cr", book_bytes(CR_HEADER, [row]), 43, 2026)

    def test_blank_rows_are_ignored_but_empty_base_is_rejected(self):
        result = generate("cr", book_bytes(CR_HEADER, [[None] * 9, cr_row()]), 43, 2026)
        self.assertEqual(result.total, 1)
        with self.assertRaisesRegex(StructureError, "no contiene rutas"):
            generate("cr", book_bytes(CR_HEADER, []), 43, 2026)

    def test_invalid_excel_fails_cleanly(self):
        with self.assertRaises(StructureError):
            generate("cr", b"not an xlsx", 43, 2026)

    def test_invalid_iso_week_fails(self):
        with self.assertRaisesRegex(StructureError, "semana indicada no existe"):
            generate("cr", book_bytes(CR_HEADER, [cr_row()]), 53, 2025)

    def test_missing_previous_sheet_fails_instead_of_reporting_false_additions(self):
        with self.assertRaisesRegex(StructureError, "pestaña"):
            generate("cr", book_bytes(CR_HEADER, [cr_row()]), 43, 2026, book_bytes(CR_HEADER, [cr_row()]))

    def test_additions_removals_and_multiple_fields_in_one_modified_record(self):
        previous = generate("cr", book_bytes(CR_HEADER, [cr_row("A"), cr_row("B")]), 42, 2026)
        current = generate("cr", book_bytes(CR_HEADER, [cr_row("A", sub="Nueva", zone="Nueva zona"), cr_row("C")]), 43, 2026, previous.content)
        row = current.summary[0]
        self.assertEqual((row["Altas"], row["Bajas"], row["Modificados"]), (1, 1, 1))
        self.assertEqual(len([c for c in current.changes if c["Tipo"] == "Cambio"]), 2)

    def test_renamed_key_is_addition_and_removal(self):
        previous = generate("vales", book_bytes(VALES_HEADER, [vales_row("Nombre anterior")]), 42, 2026)
        current = generate("vales", book_bytes(VALES_HEADER, [vales_row("Nuevo nombre")]), 43, 2026, previous.content)
        summary = next(row for row in current.summary if row["Hoja"] == "Vale Amigo")
        self.assertEqual((summary["Altas"], summary["Bajas"], summary["Modificados"]), (1, 1, 0))

    def test_previous_visual_style_and_header_text_are_preserved(self):
        previous = generate("cr", book_bytes(CR_HEADER, [cr_row()]), 42, 2026)
        book = load_workbook(BytesIO(previous.content))
        sheet = book["MEXICO"]
        sheet["D1"] = "Zona "
        sheet["A2"].font = Font(name="Arial", size=12, color="112233")
        sheet.column_dimensions["A"].width = 44
        buffer = BytesIO()
        book.save(buffer)
        book.close()
        result = generate("cr", book_bytes(CR_HEADER, [cr_row()]), 43, 2026, buffer.getvalue())
        final = load_workbook(BytesIO(result.content))
        self.assertEqual(final["MEXICO"]["D1"].value, "Zona ")
        self.assertEqual(final["MEXICO"]["A2"].font.name, "Arial")
        self.assertEqual(final["MEXICO"].column_dimensions["A"].width, 44)
        final.close()

    def test_literal_formula_text_and_leading_zero_identifiers(self):
        result = generate("cr", book_bytes(CR_HEADER, [cr_row("001"), cr_row("=2+2")]), 43, 2026)
        book = load_workbook(BytesIO(result.content), data_only=False)
        self.assertEqual(book["MEXICO"]["A2"].value, "001")
        self.assertEqual(book["MEXICO"]["A3"].value, "=2+2")
        self.assertEqual(book["MEXICO"]["A3"].data_type, "s")
        self.assertFalse(any(c.data_type == "f" for sheet in book for row in sheet for c in row))
        book.close()
        with ZipFile(BytesIO(result.content)) as archive:
            self.assertFalse(any(name.startswith("xl/externalLinks/") for name in archive.namelist()))

    def test_bundle_contains_both_correctly_named_workbooks(self):
        cr = generate("cr", book_bytes(CR_HEADER, [cr_row()]), 43, 2026)
        vales = generate("vales", book_bytes(VALES_HEADER, [vales_row()]), 43, 2026)
        with ZipFile(BytesIO(bundle([cr, vales]))) as archive:
            self.assertEqual(set(archive.namelist()), {cr.filename, vales.filename})
            self.assertEqual(archive.read(cr.filename), cr.content)


if __name__ == "__main__":
    unittest.main()
