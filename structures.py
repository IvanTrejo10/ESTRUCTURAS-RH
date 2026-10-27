"""Reglas y generación de las estructuras semanales de RH, sin servicios externos."""
from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date
from io import BytesIO, StringIO
import csv
import unicodedata
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_EXPANDED_BYTES = 120 * 1024 * 1024
COUNTRIES = (
    ("MEXICO", "México"), ("GUATEMALA", "Guatemala"), ("PERU", "Peru"),
    ("EL SALVADOR", "El Salvador"), ("COLOMBIA", "Colombia"),
    ("HONDURAS", "Honduras"), ("NICARAGUA", "Nicaragua"),
)
BRANDS = (("Rapi", "RapiVale"), ("Peru", "Vale Perú"),
          ("Vale Amigo", "Vale Amigo"), ("Viva Vale", "Viva Vale"))
CR_FIELDS = ("Ruta", "Territorio", "Subdireccion", "Zona", "Sucursal", "Marca", "País")
LATAM_FIELDS = ("Ruta", "Territorio", "Zona", "Sucursal", "Marca", "País")
VALES_FIELDS = ("Coordinacion", "Subdirección", "Zona", "Sucursal", "Marca")


class StructureError(ValueError):
    """Problema de datos que la persona puede resolver antes de descargar."""


@dataclass
class Result:
    kind: str
    week: int
    year: int
    filename: str
    content: bytes
    sheets: dict[str, list[list[str]]]
    headers: dict[str, list[str]]
    summary: list[dict]
    changes: list[dict]
    compared: bool

    @property
    def total(self) -> int:
        return sum(len(rows) for rows in self.sheets.values())


def normalize(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).casefold().split())


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _open(content: bytes, *, read_only=True):
    if not content or len(content) > MAX_FILE_BYTES:
        raise StructureError("Sube un archivo Excel .xlsx de hasta 25 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if "xl/workbook.xml" not in archive.namelist():
                raise StructureError("El archivo no es un libro Excel .xlsx válido.")
            if sum(info.file_size for info in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise StructureError("El libro es demasiado grande para procesarlo en esta página.")
        return load_workbook(BytesIO(content), read_only=read_only, data_only=True, keep_links=False)
    except StructureError:
        raise
    except Exception as exc:
        # Los errores del analizador XML varían entre versiones compatibles de openpyxl.
        raise StructureError("No se pudo leer el Excel. Revisa que sea .xlsx y que abra correctamente.") from exc


def _fields(kind: str, name: str) -> tuple[str, ...]:
    return VALES_FIELDS if kind == "vales" else CR_FIELDS if name == "MEXICO" else LATAM_FIELDS


def _groups(kind: str):
    if kind not in ("cr", "vales"):
        raise StructureError("Elige una base CR o una base de Vales.")
    return COUNTRIES if kind == "cr" else BRANDS


def _find_source(book, kind: str):
    required = CR_FIELDS if kind == "cr" else VALES_FIELDS
    wanted = {normalize(field) for field in required}
    preferred = "estructuras" if kind == "cr" else "hoja 1"
    candidates = []
    for sheet in book:
        for row_number, row in enumerate(sheet.iter_rows(max_row=10, values_only=True), start=1):
            headers = [normalize(value) for value in row]
            if wanted.issubset(headers):
                if any(headers.count(field) != 1 for field in wanted):
                    raise StructureError(f"La hoja {sheet.title} tiene encabezados requeridos repetidos.")
                candidate = (sheet, row_number, {field: headers.index(field) for field in wanted})
                if normalize(sheet.title) == preferred:
                    return candidate
                candidates.append(candidate)
                break
    if len(candidates) == 1:
        return candidates[0]
    if len(candidates) > 1:
        raise StructureError("Hay varias hojas posibles. Usa 'Estructuras' para CR o 'Hoja 1' para Vales.")
    raise StructureError("No encontré las columnas requeridas: " + ", ".join(required) + ".")


def _read_source(content: bytes, kind: str):
    book = _open(content)
    groups = _groups(kind)
    sheets = {name: [] for name, _ in groups}
    group_names = {normalize(value): name for name, value in groups}
    required = CR_FIELDS if kind == "cr" else VALES_FIELDS
    group_field = "País" if kind == "cr" else "Marca"
    errors, keys = [], set()
    try:
        sheet, header_row, indices = _find_source(book, kind)
        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            values = {field: row[indices[normalize(field)]] if indices[normalize(field)] < len(row) else None for field in required}
            if all(_blank(value) for value in values.values()):
                continue
            missing = [field for field, value in values.items() if _blank(value)]
            if missing:
                errors.append(f"Fila {row_number}: falta {', '.join(missing)}.")
                continue
            name = group_names.get(normalize(values[group_field]))
            if name is None:
                errors.append(f"Fila {row_number}: {group_field} no contemplado: {values[group_field]}.")
                continue
            route = _text(values[required[0]])
            key = (name, route)
            if key in keys:
                errors.append(f"Fila {row_number}: clave repetida en {name}: {route}.")
                continue
            keys.add(key)
            if kind == "cr" and name != "MEXICO":
                values["Territorio"] = values["Subdireccion"]
            sheets[name].append([_text(values[field]) for field in _fields(kind, name)])
        if errors:
            detail = "\n".join(errors[:10])
            rest = f"\nY {len(errors) - 10} observaciones más." if len(errors) > 10 else ""
            raise StructureError(detail + rest + "\nSi una celda contiene una fórmula, abre y guarda la base en Excel para actualizar sus valores.")
        if not any(sheets.values()):
            raise StructureError("La base no contiene rutas o coordinaciones para generar las estructuras.")
        return sheets
    finally:
        book.close()


def _previous(content: bytes, kind: str):
    book = _open(content, read_only=False)
    sheets, headers = {}, {}
    by_name = {normalize(sheet.title): sheet for sheet in book}
    try:
        for name, _ in _groups(kind):
            if normalize(name) not in by_name:
                raise StructureError(f"El archivo anterior no contiene la pestaña {name}.")
            sheet = by_name[normalize(name)]
            expected = _fields(kind, name)
            actual = [sheet.cell(1, col).value for col in range(1, len(expected) + 1)]
            if [normalize(field) for field in actual] != [normalize(field) for field in expected]:
                raise StructureError(f"Los encabezados de {name} no coinciden con el formato anterior.")
            headers[name] = [_text(value) for value in actual]
            rows = []
            for row in sheet.iter_rows(min_row=2, max_col=len(expected), values_only=True):
                if all(_blank(value) for value in row):
                    continue
                if any(_blank(value) for value in row):
                    raise StructureError(f"El archivo anterior tiene un registro incompleto en {name}.")
                rows.append([_text(value) for value in row])
            if len(rows) != len({row[0] for row in rows}):
                raise StructureError(f"El archivo anterior tiene claves repetidas en {name}.")
            sheets[name] = rows
        return book, sheets, headers, by_name
    except Exception:
        book.close()
        raise


def _compare(sheets, previous, headers):
    summary, changes = [], []
    for name, rows in sheets.items():
        old = {row[0]: row for row in previous[name]}
        new = {row[0]: row for row in rows}
        added, removed = sorted(new.keys() - old.keys()), sorted(old.keys() - new.keys())
        modified = set()
        for key in added:
            changes.append({"Tipo": "Alta", "Hoja": name, "Ruta / Coordinación": key, "Campo": "Registro", "Antes": "", "Ahora": "; ".join(f"{h}: {v}" for h, v in zip(headers[name][1:], new[key][1:]))})
        for key in removed:
            changes.append({"Tipo": "Baja", "Hoja": name, "Ruta / Coordinación": key, "Campo": "Registro", "Antes": "; ".join(f"{h}: {v}" for h, v in zip(headers[name][1:], old[key][1:])), "Ahora": ""})
        for key in sorted(new.keys() & old.keys()):
            for column, (before, after) in enumerate(zip(old[key], new[key])):
                if before != after:
                    modified.add(key)
                    changes.append({"Tipo": "Cambio", "Hoja": name, "Ruta / Coordinación": key, "Campo": headers[name][column].strip(), "Antes": before, "Ahora": after})
        summary.append({"Hoja": name, "Registros": len(rows), "Altas": len(added), "Bajas": len(removed), "Modificados": len(modified)})
    return summary, changes


def _excel(kind, sheets, headers, previous_styles=None):
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        template = previous_styles.get(normalize(name)) if previous_styles else None
        columns = len(headers[name])
        widths = (31, 18, 23, 23, 48, 21, 16) if name == "MEXICO" else (31, 23, 23, 48, 21, 16) if kind == "cr" else (36, 24, 28, 32, 20)
        for col, width in enumerate(widths, start=1):
            letter = get_column_letter(col)
            sheet.column_dimensions[letter].width = template.column_dimensions[letter].width if template else width
        sheet.sheet_view.showGridLines = template.sheet_view.showGridLines if template else True
        sheet.freeze_panes = template.freeze_panes if template else "A2"
        sheet.sheet_properties.tabColor = copy(template.sheet_properties.tabColor) if template else None
        for row_number, values in enumerate([headers[name], *rows], start=1):
            if template:
                style_row = row_number if row_number <= template.max_row else 2
                sheet.row_dimensions[row_number].height = template.row_dimensions[style_row].height
            else:
                sheet.row_dimensions[row_number].height = 19 if kind == "vales" else 15
            for col_number, value in enumerate(values, start=1):
                cell = sheet.cell(row_number, col_number, value)
                cell.data_type = "s"  # Las rutas y etiquetas son texto literal, incluso si inician con '='.
                if template:
                    source = template.cell(style_row, col_number)
                    cell.font, cell.fill, cell.border = copy(source.font), copy(source.fill), copy(source.border)
                    cell.alignment, cell.protection = copy(source.alignment), copy(source.protection)
                    cell.number_format = source.number_format
                else:
                    header = row_number == 1
                    cell.font = Font(name="Century Gothic" if kind == "cr" else "Aptos Narrow", size=8 if kind == "cr" else 11, bold=kind == "vales", color="FFFFFF" if header else "172B3A")
                    cell.alignment = Alignment(horizontal="left" if header or kind == "cr" else "center", vertical="center")
                    if header:
                        cell.fill = PatternFill("solid", fgColor="0070C0" if kind == "cr" else "0D3340")
                    if kind == "cr":
                        side = Side(style="thin", color="202020")
                        cell.border = Border(left=side, right=side, top=side, bottom=side)
        if not template or template.auto_filter.ref:
            sheet.auto_filter.ref = f"A1:{get_column_letter(columns)}{len(rows) + 1}"
    book.calculation.fullCalcOnLoad = False
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def generate(kind: str, content: bytes, week: int, year: int, previous: bytes | None = None) -> Result:
    try:
        date.fromisocalendar(year, week, 1)
    except ValueError as exc:
        raise StructureError("La semana indicada no existe en ese año. Revisa año y semana.") from exc
    sheets = _read_source(content, kind)
    previous_book = None
    styles = None
    headers = {name: list(_fields(kind, name)) for name in sheets}
    summary = [{"Hoja": name, "Registros": len(rows)} for name, rows in sheets.items()]
    changes = []
    try:
        if previous is not None:
            previous_book, old, headers, styles = _previous(previous, kind)
            summary, changes = _compare(sheets, old, headers)
        output = _excel(kind, sheets, headers, styles)
    finally:
        if previous_book:
            previous_book.close()
    filename = f"Estructura {'por país' if kind == 'cr' else 'Vales'} - SEM {week}.xlsx"
    return Result(kind, week, year, filename, output, sheets, headers, summary, changes, previous is not None)


def changes_csv(result: Result) -> bytes:
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=("Tipo", "Hoja", "Ruta / Coordinación", "Campo", "Antes", "Ahora"))
    writer.writeheader()
    for row in result.changes:
        writer.writerow({key: "'" + value if value.startswith(("=", "+", "-", "@")) else value for key, value in row.items()})
    return output.getvalue().encode("utf-8-sig")


def bundle(results: list[Result]) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        for result in results:
            archive.writestr(result.filename, result.content)
    return buffer.getvalue()
